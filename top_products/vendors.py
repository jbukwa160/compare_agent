"""EU vendors' price lists from the company's "EU vendors" Google Drive folder (or its .zip download)."""
from __future__ import annotations

import gzip
import hashlib
import json
import os
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from compare_agent.config import load_env

from . import drive
from .config import CACHE, ROOT
from .pages import norm_gtin

PARSER_VERSION = 2      # bump when _read() changes, so cached results are read again


def _drive_for_desktop(rel: str) -> Path | None:
    """The folder as Google Drive for desktop shows it: G:\\Shared drives\\... (any letter), or the macOS location."""
    rel_path = Path(rel.replace("\\", "/"))
    if rel_path.is_absolute():
        return rel_path if rel_path.is_dir() else None
    bases = [Path(f"{d}:/") for d in "GHIJKLMNOPQRSTUVWXYZDEF"] if os.name == "nt" else []
    bases += sorted((Path.home() / "Library" / "CloudStorage").glob("GoogleDrive-*"))
    bases += [Path.home() / "Google Drive", Path("/mnt/g")]
    for base in bases:
        try:
            cand = base / rel_path
            if cand.is_dir():
                return cand
            if cand.parent.is_dir():          # "EU vendor" / "EU vendors"
                stem = cand.name.lower().rstrip("s")
                hits = [c for c in cand.parent.iterdir() if c.is_dir() and c.name.lower().rstrip("s") == stem]
                if hits:
                    return hits[0]
        except OSError:                       # an empty card reader / disconnected drive letter
            continue
    return None


def find_source(arg: str | None, cfg: dict) -> tuple[Path | None, str]:
    """(folder or .zip, where it came from). The first that exists wins (see settings.yaml -> vendors)."""
    if arg:
        return Path(arg).expanduser(), "зададен път"
    load_env()
    key = os.environ.get("GOOGLE_SERVICE_ACCOUNT_FILE", "").strip()
    if key and cfg.get("drive_folder_id"):
        out = CACHE / "drive" / cfg["drive_folder_id"]
        return drive.sync(drive.session(key), cfg["drive_folder_id"], out), "Google Drive (API)"
    synced = _drive_for_desktop(cfg.get("drive_path", ""))
    if synced:
        return synced, "Google Drive"
    found = []
    for folder in (ROOT, ROOT / "input", ROOT.parent, Path.home() / "Downloads"):
        found += [p for p in folder.glob("EU vendors*") if p.suffix == ".zip" or p.is_dir()]
    if found:
        return max(found, key=lambda p: p.stat().st_mtime), "изтеглен архив от Google Drive"
    return None, ""


def _extract(zip_path: Path) -> Path:
    """Unzips the spreadsheets once into the cache (paths checked: nothing lands outside it)."""
    st = zip_path.stat()
    out = CACHE / "vendors" / f"{zip_path.stem}-{st.st_size}-{int(st.st_mtime)}"
    if not (out / ".done").exists():
        out.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(zip_path) as zf:
            for info in zf.infolist():
                target = (out / info.filename).resolve()
                if info.is_dir() or not info.filename.lower().endswith(".xlsx") or out.resolve() not in target.parents:
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(zf.read(info))
        (out / ".done").write_text("ok")
    tops = [p for p in out.iterdir() if p.is_dir()]
    return tops[0] if len(tops) == 1 else out       # the zip holds one "EU vendors" folder


def _rows(path: Path, sheet: str | None = None):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = [str(h).strip() if h is not None else "" for h in next(it, [])]
        for r in it:
            yield dict(zip(head, r))
    finally:
        wb.close()


def _num(v) -> float | None:
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "") else None
    except ValueError:
        return None


def _read(root: Path) -> tuple[list[dict], list[dict]]:
    """(price rows, one line per vendor folder). A vendor's prices: supplier_products.xlsx plus the
    supplier_products sheets of its offer files; products.xlsx names the barcodes."""
    items, vendors = [], []
    for vendor_dir in sorted(p for p in root.iterdir() if p.is_dir()):
        names = {}
        if (vendor_dir / "products.xlsx").exists():
            for r in _rows(vendor_dir / "products.xlsx"):
                e = norm_gtin(r.get("ean_main"))
                if e:
                    names[e] = (str(r.get("name") or "").strip(), str(r.get("brand") or "").strip())
        sources = [("supplier_products.xlsx", None)] if (vendor_dir / "supplier_products.xlsx").exists() else []
        for f in sorted(vendor_dir.glob("*.xlsx")):
            if f.name in ("supplier_products.xlsx", "products.xlsx") or f.name.startswith("~$"):
                continue
            try:
                wb = load_workbook(f, read_only=True)
                sheets = wb.sheetnames
                wb.close()
            except Exception:
                continue
            sources += [(f.name, sh) for sh in sheets if sh.startswith("supplier_products")]
        before, raw, no_price, no_ean = len(items), 0, 0, 0
        for fname, sheet in sources:
            for r in _rows(vendor_dir / fname, sheet):
                if not any(v not in (None, "") for v in r.values()):
                    continue
                raw += 1
                e = norm_gtin(r.get("supplier_ean")) or norm_gtin(r.get("supplier_sku"))
                price = _num(r.get("price"))
                no_ean += not e
                no_price += bool(e) and (price is None or price <= 0)
                if not e or price is None or price <= 0:
                    continue
                name, brand = names.get(e, ("", ""))
                items.append({"vendor": vendor_dir.name, "ean": e, "name": str(r.get("name") or "").strip() or name,
                              "brand": brand, "price": price, "currency": str(r.get("currency") or "EUR").strip().upper(),
                              "stock": r.get("stock"), "source": fname})
        mine = items[before:]
        note = None
        if not sources:
            note = "няма ценова листа (supplier_products) в папката"
        elif not raw:
            note = "ценовата листа (supplier_products) е празна"
        elif not mine:
            note = f"ценовата листа няма цени ({raw} реда)" if no_price else f"ценовата листа няма баркодове ({raw} реда)"
        elif no_price or no_ean:
            note = f"{no_price + no_ean} реда без цена или баркод не са взети"
        vendors.append({"vendor": vendor_dir.name, "rows": len(mine), "eans": len({i["ean"] for i in mine}), "note": note})
    return items, vendors


def _fingerprint(source: Path) -> str:
    """Changes whenever any spreadsheet in the folder (or the .zip) changes."""
    files = [source] if source.is_file() else sorted(source.rglob("*.xlsx"))
    h = hashlib.sha1()
    for f in files:
        st = f.stat()
        h.update(f"{f.relative_to(source) if source.is_dir() else f.name}|{st.st_size}|{st.st_mtime_ns}".encode())
    return h.hexdigest()[:16]


def load(source: Path) -> tuple[list[dict], list[dict]]:
    """(price rows, vendors). Parsed once per version of the folder / .zip, then read from the cache."""
    cache = CACHE / "vendors" / f"items-v{PARSER_VERSION}-{_fingerprint(source)}.json.gz"
    if cache.exists():
        data = json.loads(gzip.decompress(cache.read_bytes()))
    else:
        items, vendors = _read(_extract(source) if source.suffix == ".zip" else source)
        data = {"items": items, "vendors": vendors}
        cache.parent.mkdir(parents=True, exist_ok=True)
        cache.write_bytes(gzip.compress(json.dumps(data, ensure_ascii=False).encode()))
    return data["items"], data["vendors"]
