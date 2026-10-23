"""Reads the traffic data back from the Excel and writes the whole workbook."""
from __future__ import annotations

import os
from pathlib import Path
from urllib.parse import unquote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CATEGORIES = {"Детски": "Детски магазини", "Аптеки": "Аптеки", "Технологии": "Технологии",
              "Хранителни добавки": "Хранителни добавки"}
PRODUCT_COLS = [("Магазин", 18, "store"), ("Водещ", 8, "lead"), ("№", 5, "rank"), ("Продукт", 58, "name"),
                ("Марка", 16, "brand"), ("Категория в магазина", 34, "crumbs"), ("Трафик/месец (оценка)", 12, "etv"),
                ("Ключови думи (брой)", 11, "keywords"), ("Основна ключова дума", 30, "kw"),
                ("Търсения/месец на ключовата дума", 13, "kw_volume"), ("Позиция в Google", 10, "kw_rank"),
                ("Цена от (€)", 10, "price_min"), ("до (€)", 9, "price_max"), ("Наличност", 13, "availability"),
                ("Линк", 60, "url"), ("Бележка", 40, "note")]
VENDOR_COLS = [("Скор (1–10)", 8, "score"), ("Продукт", 46, "name"), ("Баркод (EAN)", 15, "ean"), ("Доставчик", 16, "vendor"),
               ("Цена доставчик (€)", 11, "vendor_price"), ("Цена доставчик +20% (€)", 11, "vendor_price_vat"),
               ("Най-ниска цена в БГ (€)", 11, "bg_price"), ("Магазин с най-ниска цена", 14, "bg_shop"),
               ("Печалба на брой (€)", 10, "profit"), ("Марж", 8, "margin"), ("Трафик/месец", 10, "traffic"),
               ("Наличност при доставчика (бр.)", 11, "stock"), ("Други доставчици (€)", 26, "other_vendors"),
               ("БГ магазини от таблицата (€)", 40, "shops"), ("Как е намерен", 14, "how"), ("Линк", 50, "bg_url")]
LEGACY_META = {"data_date": "2026-10-17", "checked": "21–22.10.2026", "pages": "1000 / 300"}

BOLD, HEAD_FONT = Font(bold=True), Font(bold=True, color="FFFFFF")
HEAD_FILL, BAND_FILL = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="EAF1FB")
SCORE_FILLS = [(7, PatternFill("solid", fgColor="C6EFCE")), (4, PatternFill("solid", fgColor="FFEB9C")),
               (1, PatternFill("solid", fgColor="EDEDED"))]
THIN = Border(bottom=Side(style="thin", color="D0D7E2"))
WRAP, TOP = Alignment(wrap_text=True, vertical="top"), Alignment(vertical="top")
LINK = Font(color="0563C1", underline="single")


def find(folder: Path, default_name: str) -> Path:
    """The newest "Топ продукти по трафик*.xlsx" in the folder, else a new file of that name."""
    found = [p for p in folder.glob("Топ продукти по трафик*.xlsx") if not p.name.startswith("~$")]
    return max(found, key=lambda p: p.stat().st_mtime) if found else folder / default_name


def read(path: Path) -> tuple[list[dict], list[dict], dict]:
    """(shops, product rows, meta) from a workbook this tool (or the first runs) wrote."""
    wb = load_workbook(path)
    ws = wb["Обобщение"]
    head_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Категория")
    col = {ws.cell(head_row, c).value: c for c in range(1, ws.max_column + 1)}
    shops = []
    for r in range(head_row + 1, ws.max_row + 1):
        g = lambda name: ws.cell(r, col[name]).value if name in col else None  # noqa: E731
        if g("Категория"):
            traffic = g("Органичен трафик на сайта/месец")
            shops.append({"category": g("Категория"), "store": g("Магазин"), "domain": g("Сайт"),
                          "traffic": traffic if isinstance(traffic, (int, float)) else None,
                          "lead": (g("Водещ (топ 6)") or g("В таблицата")) == "да", "note": g("Бележка")})
    ws = wb["Всички продукти"]
    keys = {h: k for h, _, k in PRODUCT_COLS} | {"Категория": "category", "№ в магазина": "rank"}
    col = {keys[ws.cell(4, c).value]: c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value in keys}
    rows = []
    for r in range(5, ws.max_row + 1):
        row = {k: ws.cell(r, c).value for k, c in col.items()}
        link = ws.cell(r, col["url"]).hyperlink
        row["url"] = link.target if link is not None else row["url"]
        row.setdefault("lead", "да")
        rows.append(row)
    meta = dict(LEGACY_META)
    if "_meta" in wb.sheetnames:
        meta.update({k: v for k, v in wb["_meta"].iter_rows(values_only=True) if k})
    return shops, rows, meta


def _header(ws, row: int, cols: list) -> None:
    for c, (h, w, _) in enumerate(cols, 1):
        cell = ws.cell(row, c, h)
        cell.font, cell.fill, cell.alignment = HEAD_FONT, HEAD_FILL, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 45


def _notes(ws, first_row: int, notes: list[str], width: int) -> int:
    for i, n in enumerate(notes, start=first_row):
        ws.cell(i, 1, n).alignment = WRAP
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=width)
        ws.row_dimensions[i].height = 15 if len(n) < 110 else 32 if len(n) < 260 else 46
    return first_row + len(notes)


def _summary(wb: Workbook, shops: list[dict], rows: list[dict], meta: dict, vendor_line: str) -> None:
    ws = wb.active
    ws.title = "Обобщение"
    ws["A1"] = "Най-посещавани продукти в онлайн магазините в България"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Детски магазини, аптеки, технологии, хранителни добавки · проверено на {meta['checked']} · "
                f"източник: DataForSEO Labs (Google, България, български език, данни към {meta['data_date']})")
    pages = str(meta["pages"]).split(" / ")
    notes = [
        "Как да се чете таблицата:",
        "• Трафик/месец = оценка на месечните посещения на страницата на продукта от органичното (неплатено) търсене "
        "в Google в България. Изчислява се от позициите на страницата в Google и броя търсения на всяка ключова дума.",
        "• Не включва платена реклама, директни посещения, социални мрежи и търсенето вътре в сайта. Числата са "
        "оценки: подходящи за сравнение между продукти и магазини, не са точен брой посещения.",
        "• Водещи магазини (колона „Водещ“ = „да“) = първите 6 магазина по органичен трафик от Google в България във всяка "
        "категория. Сайтове, чийто трафик идва почти изцяло от статии, не са сред тях. "
        "Всички останали проверени магазини също са в таблицата, с „не“ в колона „Водещ“.",
        f"• До 20 продукта на магазин; при по-малките магазини може да са по-малко. За водещите магазини са прегледани "
        f"първите {pages[0]} страници на сайта по трафик, за останалите първите {pages[-1]}. "
        "При eMAG и Ozone са взети само продуктите от техника, при Jumbo само играчките и бебешките стоки.",
        "• Само продуктови страници (без категории, марки и статии). Всяка страница е отворена и проверена; цената и "
        "наличността са от сайта към датата на проверката. Един продукт с няколко адреса е събран в един ред.",
        "• Основна ключова дума = търсенето в Google, което носи най-много посещения на страницата; "
        "Търсения/месец = колко пъти месечно се търси тя в България; Позиция = мястото на страницата в Google за нея.",
    ] + ([f"• Лист „EU доставчици“: {vendor_line}"] if vendor_line else [])
    r = _notes(ws, 4, notes, 9) + 1
    ws.cell(4, 1).font = BOLD
    count = {}
    for row in rows:
        count[row["store"]] = count.get(row["store"], 0) + 1
    head = ["Категория", "Магазин", "Сайт", "Органичен трафик на сайта/месец", "Водещ (топ 6)", "Продукти в таблицата",
            "Бележка", "Най-посещаван продукт", "Трафик/месец на продукта"]
    _header(ws, r, [(h, w, None) for h, w in zip(head, (19, 18, 20, 15, 10, 11, 34, 55, 13))])
    for cat in CATEGORIES:
        for shop in sorted((s for s in shops if s["category"] == cat), key=lambda s: (not s["lead"], -(s["traffic"] or 0))):
            r += 1
            first = min((x for x in rows if x["store"] == shop["store"]), key=lambda x: x["rank"], default=None)
            vals = [cat, shop["store"], shop["domain"], round(shop["traffic"]) if shop["traffic"] else "няма данни",
                    "да" if shop["lead"] else "не", count.get(shop["store"], 0), shop["note"] or None,
                    first["name"] if first else None, first["etv"] if first else None]
            for c, v in enumerate(vals, 1):
                cell = ws.cell(r, c, v)
                cell.alignment = WRAP if c in (7, 8) else TOP
                if shop["lead"]:
                    cell.font = BOLD
            ws.cell(r, 4).number_format = ws.cell(r, 9).number_format = "#,##0"


def _products(wb: Workbook, title: str, data: list[dict], with_category: bool, scope: str) -> None:
    ws = wb.create_sheet(title)
    cols = ([("Категория", 17, "category")] + [("№ в магазина", 8, "rank") if c[2] == "rank" else c for c in PRODUCT_COLS]
            if with_category else PRODUCT_COLS)
    ws.cell(1, 1, f"{title}: продуктите с най-много трафик от Google в България").font = Font(bold=True, size=13)
    ws.cell(2, 1, "Трафик/месец = оценка на посещенията от органичното търсене в Google; подробности в лист „Обобщение“. "
            + ("" if with_category else "Първо са 6-те водещи магазина, след тях останалите проверени. ") + scope)
    _header(ws, 4, cols)
    band, prev = False, None
    for i, row in enumerate(data, start=5):
        if row["store"] != prev and not with_category:
            band, prev = not band, row["store"]
        for c, (_, _, key) in enumerate(cols, 1):
            v = row.get(key)
            cell = ws.cell(i, c, v if v not in ("", None) else None)
            cell.alignment = WRAP if key in ("name", "crumbs", "note", "kw") else TOP
            cell.border = THIN
            if band and not with_category:
                cell.fill = BAND_FILL
            if key in ("etv", "keywords", "kw_volume"):
                cell.number_format = "#,##0"
            elif key in ("price_min", "price_max"):
                cell.number_format = "#,##0.00"
            elif key == "url" and v:
                cell.value, cell.hyperlink, cell.font = unquote(v), v, LINK
    ws.freeze_panes = ws.cell(5, 5 if with_category else 4)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{max(5, 4 + len(data))}"


def _vendors(wb: Workbook, table: list[dict], left_out: dict, source: str, checked: str, score_cfg: dict, vat: float) -> None:
    ws = wb.create_sheet("EU доставчици", 1)
    ws["A1"] = "Продукти от таблицата, които EU доставчиците предлагат по-евтино от българските магазини"
    ws["A1"].font = Font(bold=True, size=13)
    pe, mg, tr = score_cfg["profit_eur"], score_cfg["margin"], score_cfg["traffic"]
    fmt = lambda xs: " / ".join(f"{x:,}".replace(",", " ") for x in xs)  # noqa: E731
    _notes(ws, 2, [
        f"Цена доставчик = най-ниската цена при EU доставчик, без ДДС, в EUR. +{vat:.0%} = цена доставчик × "
        + f"{1 + vat:g}".replace(".", ",") + ". Печалба на брой = най-ниска цена в български магазин (с ДДС) − "
        f"цена доставчик +{vat:.0%}. Марж = печалба / цена в БГ.",
        f"Скор 1–10 = точки за печалба на брой ({fmt(pe)} €), марж ({' / '.join(f'{m:.0%}' for m in mg)}) и трафик "
        f"({fmt(tr)} посещения/месец). Под {score_cfg['min_profit_eur']} € печалба на брой = 1 (не си струва вноса); "
        f"марж под {score_cfg['low_margin']:.0%} = най-много {score_cfg['low_margin_max_score']}. "
        f"Махнати: {left_out['cheaper_in_bg']} продукта, които в България са по-евтини, и {left_out['no_bg_price']} без цена в БГ.",
        f"Доставчици: {source} · проверено на {checked}. „По име – за проверка“ = съвпадение по име, което никой не е "
        "проверил: прегледайте го и го добавете в top_products/review.yaml.",
    ], len(VENDOR_COLS))
    _header(ws, 5, VENDOR_COLS)
    for i, row in enumerate(table, start=6):
        for c, (_, _, key) in enumerate(VENDOR_COLS, 1):
            v = row.get(key)
            if key == "stock":
                v = int(float(v)) if str(v or "").replace(".", "", 1).isdigit() else (v or None)
            cell = ws.cell(i, c, v if v not in ("", None) else None)
            cell.alignment = WRAP if key in ("name", "shops", "other_vendors") else TOP
            cell.border = THIN
            if key in ("vendor_price", "vendor_price_vat", "bg_price", "profit"):
                cell.number_format = "#,##0.00"
            elif key == "margin":
                cell.number_format = "0%"
            elif key in ("traffic", "stock"):
                cell.number_format = "#,##0"
            elif key == "ean":
                cell.number_format = "@"
            elif key == "score":
                cell.font, cell.alignment = Font(bold=True, size=12), Alignment(horizontal="center", vertical="top")
                cell.fill = next(f for lim, f in SCORE_FILLS if v >= lim)
            elif key == "bg_url" and v:
                cell.value, cell.hyperlink, cell.font = unquote(v), v, LINK
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(VENDOR_COLS))}{max(6, 5 + len(table))}"


def write(path: Path, shops: list[dict], rows: list[dict], meta: dict, vendors: dict | None) -> None:
    """The whole workbook; saved next to the target first, so a failed save never leaves half a file."""
    wb = Workbook()
    _summary(wb, shops, rows, meta, vendors["line"] if vendors else "")
    if vendors:
        _vendors(wb, vendors["table"], vendors["left_out"], vendors["source"], meta["checked_vendors"],
                 vendors["score_cfg"], vendors["vat"])
    order = {s["store"]: i for i, s in enumerate(sorted(shops, key=lambda s: (not s["lead"], -(s["traffic"] or 0))))}
    scope = {"Детски": "Jumbo: само играчки и бебешки стоки.", "Технологии": "eMAG и Ozone: само техника."}
    for cat, title in CATEGORIES.items():
        data = sorted((x for x in rows if x["category"] == cat), key=lambda x: (order.get(x["store"], 999), x["rank"]))
        _products(wb, title, data, with_category=False, scope=scope.get(cat, ""))
    _products(wb, "Всички продукти", sorted(rows, key=lambda x: -(x["etv"] or 0)), with_category=True, scope="")
    ws = wb.create_sheet("_meta")
    for k in ("data_date", "checked", "pages"):
        ws.append([k, meta.get(k)])
    ws.sheet_state = "hidden"
    tmp = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, path)
