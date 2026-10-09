"""The Excel report: laid out like the input file, odd characters, and a locked output file."""
import openpyxl
from openpyxl import Workbook

from compare_agent.config import load_settings
from compare_agent.models import Evaluation, Offer, Product, Traffic
from compare_agent.report import write_json, write_report
from compare_agent.scoring import evaluate, market

SETTINGS = load_settings()


def evaluations():
    out = []
    for row, (name, price, qty) in enumerate([("Samsung 990 PRO 1TB", 150, 40),
                                              ("Weird «name» 🎧 with \x0b control char", 12, None),
                                              ("Not sold anywhere", None, None)], start=2):
        offers = []
        if row < 4:
            for site, scope, net, stock in (("Technopolis", "local", 199.9, True), ("Klarna DE", "global", 170, True)):
                o = Offer(site=site, scope=scope, title=f"{name} listing\x07", url="https://x/y", price=net * 1.2,
                          in_stock=stock)
                o.price_eur_net, o.match_score, o.match_reason = net, 1.0, "EAN match"
                offers.append(o)
        p = Product(row=row, name=name, ean="8806094215021", supplier_price=price, order_qty=qty,
                    raw={"internal_sku": None, "name": name, "ean_main": "8806094215021", "description": None,
                         "sale_price": price})
        ev = Evaluation(product=p, local=market("local", offers), global_=market("global", offers), offers=offers,
                        errors=["Some site: HTTP 403"],
                        traffic=[Traffic(database="Bulgaria", keyword="x", volume=90, source="DataForSEO",
                                         scope="local")])
        evaluate(ev, SETTINGS)
        out.append(ev)
    return out


def test_report_looks_like_the_input_file_plus_result_columns(tmp_path):
    path = write_report(evaluations(), SETTINGS, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Products", "Shop prices"]
    ws = wb["Products"]
    headers = [c.value for c in ws[1]]
    # input columns first, in the input order, without the ones empty for every product ("description")
    assert headers[:3] == ["name", "ean_main", "sale_price"]
    assert "description" not in headers
    for h in ("We sell it", "Cheapest in Bulgaria", "Margin in Bulgaria", "Searches / month", "Last 12 months", "Sells",
              "Score 1-10", "Verdict"):
        assert h in headers
    vat_col = headers.index("Supplier price + VAT (20%)") + 1
    assert ws.cell(3, vat_col).value == 180.0          # 150 supplier price + 20 % VAT
    assert headers.index("Supplier price + VAT (20%)") == headers.index("Supplier price") + 1
    names = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1)]
    assert names[0] == "Samsung 990 PRO 1TB" and len(names) == 3   # one row per product, input order
    assert ws.cell(3, 2).value == "8806094215021"                   # EAN kept as text
    assert write_json(evaluations(), tmp_path / "r.json").exists()


def test_margin_in_bulgaria_and_abroad_side_by_side(tmp_path):
    ws = openpyxl.load_workbook(write_report(evaluations(), SETTINGS, tmp_path / "r.xlsx"))["Products"]
    headers = [c.value for c in ws[1]]
    bg, abroad = headers.index("Margin in Bulgaria") + 1, headers.index("Margin abroad") + 1
    assert abroad == bg + 1
    assert ws.cell(3, bg).value == round((199.9 - 150) / 199.9, 4)   # cheapest BG shop, without VAT
    assert ws.cell(3, abroad).value == round((170 - 150) / 170, 4)   # cheapest worldwide shop, without VAT
    assert ws.cell(4, abroad).value is not None and ws.cell(5, abroad).value is None   # no supplier price
    local = write_report(evaluations(), SETTINGS, tmp_path / "bg.xlsx", meta={"scope": "local"})
    assert "Margin abroad" not in [c.value for c in openpyxl.load_workbook(local)["Products"][1]]


def test_locked_report_is_saved_under_a_new_name(tmp_path, monkeypatch):
    real_save, calls = Workbook.save, []

    def save(self, filename):
        calls.append(filename)
        if len(calls) == 1:
            raise PermissionError("file is open in Excel")
        real_save(self, filename)

    monkeypatch.setattr(Workbook, "save", save)
    path = write_report(evaluations(), SETTINGS, tmp_path / "r.xlsx")
    assert path != tmp_path / "r.xlsx" and path.exists()


def test_selling_rows_are_marked_and_bulgaria_run_has_no_worldwide_column(tmp_path):
    from compare_agent.catalogue import assign_groups
    evs = evaluations()
    evs[0].product.sku = "Dilon-146548"
    evs[0].product.raw["internal_sku"] = "Dilon-146548"
    assign_groups([e.product for e in evs])
    ws = openpyxl.load_workbook(write_report(evs, SETTINGS, tmp_path / "bg.xlsx", meta={"scope": "local"}))["Products"]
    headers = [c.value for c in ws[1]]
    assert "Cheapest worldwide" not in headers
    sells_col = headers.index("We sell it") + 1
    assert ws.cell(3, sells_col).value == "✔ Yes" and ws.cell(4, sells_col).value == "No"
    assert ws.cell(3, 1).value == "Dilon-146548"
    full = openpyxl.load_workbook(write_report(evaluations(), SETTINGS, tmp_path / "all.xlsx"))["Products"]
    assert "Cheapest worldwide" in [c.value for c in full[1]]
